from pathlib import Path

import pytest

from src.modules.chat.application.knowledge.loaders import (
    load_document,
    load_docx,
    load_text_file,
    load_xlsx,
)


def test_load_document_dinh_dang_khong_ho_tro_nem_value_error(tmp_path: Path):
    path = tmp_path / "file.xyz"
    path.write_text("x", encoding="utf-8")
    with pytest.raises(ValueError, match="không được hỗ trợ"):
        load_document(path)


def test_load_document_txt_va_md(tmp_path: Path):
    txt = tmp_path / "note.txt"
    txt.write_text(
        "\n".join(f"Nội dung điều {i}. khoản quy định số {i}" for i in range(50)),
        encoding="utf-8",
    )
    chunks = load_document(txt)
    assert chunks
    assert all(c["metadata"]["document"] == "note.txt" for c in chunks)
    assert all(c["metadata"]["document_type"] == "text" for c in chunks)

    md = tmp_path / "readme.md"
    md.write_text("# Tiêu đề\n\nNội dung", encoding="utf-8")
    assert load_document(md)


def test_load_text_file_bo_header_lap(tmp_path: Path):
    path = tmp_path / "t.txt"
    body = "\n".join(["Trường Đại học Thủy Lợi"] * 4 + ["Điều 1. Quy định học phí"])
    path.write_text(body, encoding="utf-8")
    chunks = load_text_file(path)
    joined = "\n".join(c["text"] for c in chunks)
    assert "Trường Đại học Thủy Lợi" not in joined
    assert "Điều 1. Quy định học phí" in joined


def test_load_xlsx_loc_dong_tong_va_tao_block(tmp_path: Path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "MonHoc"
    ws.append(["Mã", "Tên"])
    ws.append(["MH01", "Lập trình"])
    ws.append(["", "Tổng cộng"])
    path = tmp_path / "bang.xlsx"
    wb.save(path)

    chunks = load_xlsx(path)
    assert chunks
    joined = " ".join(b["text"] for b in chunks)
    assert "MH01" in joined
    assert "Lập trình" in joined
    assert "Tổng cộng" not in joined


def test_load_docx_chuyen_thanh_chunk(tmp_path: Path):
    from docx import Document

    doc = Document()
    doc.add_paragraph("Điều 1. Quy chế đào tạo đại học")
    doc.add_paragraph("Nội dung chi tiết của điều khoản được quy định như sau.")
    path = tmp_path / "quyche.docx"
    doc.save(path)

    chunks = load_docx(path)
    assert chunks
    assert all(c["metadata"]["document"] == "quyche.docx" for c in chunks)
    assert all(c["metadata"]["document_type"] == "docx" for c in chunks)
