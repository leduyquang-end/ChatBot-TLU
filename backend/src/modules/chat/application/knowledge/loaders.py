from __future__ import annotations

import logging
import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from src.modules.chat.application.knowledge.chunking import chunk_text
from src.modules.chat.application.knowledge.cleaning import (
    clean_text,
    strip_repeated_headers,
)
from src.modules.chat.application.knowledge.table_processing import (
    drop_total_and_note_rows,
    fill_merged_cells,
    rows_to_blocks,
)

logger = logging.getLogger(__name__)

# Word 97-2003 (.doc) không đọc được bằng python-docx → dùng Word COM (Windows)
# để chuyển sang .docx tạm rồi đọc như docx. Chạy nền, không hiện cửa sổ.
# Truyền đường dẫn qua biến môi trường (env) — PowerShell 5.1 hỏng chữ Unicode nếu
# truyền qua tham số lệnh.
_CONVERT_DOC_PS = r"""
$Src = $env:TLU_DOC_SRC
$Dst = $env:TLU_DOC_DST
$ErrorActionPreference = 'Stop'
Unblock-File -Path $Src -ErrorAction SilentlyContinue
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
try {
    $doc = $word.Documents.Open($Src, $false, $true, $false)
    $doc.SaveAs2($Dst, 16)
    $doc.Close([ref]$false)
} finally {
    $word.Quit()
}
"""


def _chunks_from_docx(
    docx_path: Path, *, document_type: str, source_name: str
) -> list[dict]:
    from docx import Document

    doc = Document(str(docx_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = strip_repeated_headers(clean_text("\n".join(paragraphs)))
    return chunk_text(
        text,
        metadata={"document": source_name, "document_type": document_type},
    )


def load_docx(path: Path) -> list[dict]:
    return _chunks_from_docx(path, document_type="docx", source_name=path.name)


def _convert_doc_to_docx(src: Path, dst: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="tlu-ps-") as tmp:
        script = Path(tmp) / "convert.ps1"
        script.write_text(_CONVERT_DOC_PS, encoding="utf-8")
        env = {
            **os.environ,
            "TLU_DOC_SRC": str(src),
            "TLU_DOC_DST": str(dst),
        }
        proc = subprocess.run(
            [
                "powershell",
                "-NoProfile",
                "-ExecutionPolicy",
                "Bypass",
                "-File",
                str(script),
            ],
            env=env,
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
    if proc.returncode != 0 or not dst.exists():
        detail = (proc.stderr or proc.stdout or "")[:300]
        raise ValueError(f"Word không convert được .doc: {detail}")


def load_doc(path: Path) -> list[dict]:
    """Word 97-2003 (.doc): chuyển sang .docx tạm bằng Word COM, rồi đọc như docx.

    Không có Word (Linux/Docker) → ValueError → pipeline ghi vào documents_failed.
    """
    path = path.resolve()  # Word COM cần đường dẫn tuyệt đối
    tmp_dir = Path(tempfile.mkdtemp(prefix="tlu-doc-"))
    try:
        dst = tmp_dir / "converted.docx"
        _convert_doc_to_docx(path, dst)
        return _chunks_from_docx(dst, document_type="doc", source_name=path.name)
    except Exception as exc:
        raise ValueError(f"Không đọc được file .doc (cần Word trên Windows): {path.name}") from exc
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


def _extract_pdf_text(path: Path) -> str:
    # pypdfium2 (PDFium) giữ nguyên từ/câu tốt hơn pypdf với PDF font Việt
    # (pypdf hay vỡ "Mã học phần" thành "Mã h/ọc/phần")
    try:
        import pypdfium2 as pdfium
    except ImportError:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    doc = pdfium.PdfDocument(str(path))
    try:
        return "\n".join(
            doc[i].get_textpage().get_text_range() for i in range(len(doc))
        )
    finally:
        doc.close()


def load_pdf(path: Path) -> list[dict]:
    text = strip_repeated_headers(clean_text(_extract_pdf_text(path)))
    return chunk_text(
        text,
        metadata={"document": path.name, "document_type": "pdf"},
    )


def load_xlsx(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(str(path), data_only=True)
    blocks: list[dict] = []
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c is not None and str(c).strip() != "" for c in r)]
        rows = fill_merged_cells(rows)
        rows = drop_total_and_note_rows(rows)
        if not rows:
            continue
        headers = [str(c or "") for c in rows[0]]
        blocks.extend(
            rows_to_blocks(rows, sheet=sheet_name, headers=headers)
        )
    return blocks


def load_text_file(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    text = strip_repeated_headers(clean_text(text))
    return chunk_text(
        text,
        metadata={"document": path.name, "document_type": "text"},
    )


def load_document(path: Path) -> list[dict]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return load_docx(path)
    if suffix == ".doc":
        return load_doc(path)
    if suffix == ".pdf":
        return load_pdf(path)
    if suffix == ".xlsx":
        return load_xlsx(path)
    if suffix in {".txt", ".md"}:
        return load_text_file(path)
    raise ValueError(f"Định dạng không được hỗ trợ: {suffix}")
